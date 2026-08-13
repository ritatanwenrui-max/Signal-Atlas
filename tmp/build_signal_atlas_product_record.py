from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/tanwenrui/Documents/Codex/2026-08-11/wo-m")
OUT_DIR = ROOT / "output"
OUT_PATH = OUT_DIR / "Signal-Atlas-从0到1产品记录.docx"

ACCENT = "2E74B5"
NAVY = "17365D"
TEXT = "1F2937"
MUTED = "5F6B7A"
LIGHT = "EAF2F8"
LIGHTER = "F7F9FC"
GRID = "D8DEE8"
WHITE = "FFFFFF"
GOLD = "C9912A"
GREEN = "2F855A"
RED = "B54747"

BODY_FONT = "Hiragino Sans GB"
CJK_FONT = "Hiragino Sans GB"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_twips: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")


def set_table_width(table, width_twips=9360) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_twips))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")


def set_table_borders(table, color=GRID, size="4") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_run_font(run, size=None, bold=None, color=None, italic=None, name=BODY_FONT) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), CJK_FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def add_run(p, text, *, bold=False, color=TEXT, size=11, italic=False):
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color, italic=italic)
    return run


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_repeat_header(table) -> None:
    set_repeat_table_header(table.rows[0])


def set_cell_text(cell, text, *, bold=False, color=TEXT, size=9.5, align=None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(0)
    pf.line_spacing = 1.05
    add_run(p, text, bold=bold, color=color, size=size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    set_run_font(run, size=8.5, color=MUTED)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t")
    txt.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, txt, fld_end])


def add_hyperlink(paragraph, text: str, url: str) -> None:
    rid = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), BODY_FONT)
    fonts.set(qn("w:hAnsi"), BODY_FONT)
    fonts.set(qn("w:eastAsia"), CJK_FONT)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), ACCENT)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([fonts, color, underline])
    r.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    link.append(r)
    paragraph._p.append(link)


def add_bottom_border(paragraph, color=ACCENT, size="12", space="6") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    pbdr = p_pr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        p_pr.append(pbdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)


def add_numbering_definitions(doc: Document) -> tuple[int, int]:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    next_abs = max(abstract_ids, default=-1) + 1
    next_num = max(num_ids, default=0) + 1

    def make_abstract(abs_id: int, num_fmt: str, lvl_text: str, font: str | None = None) -> None:
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abs_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        fmt = OxmlElement("w:numFmt")
        fmt.set(qn("w:val"), num_fmt)
        text = OxmlElement("w:lvlText")
        text.set(qn("w:val"), lvl_text)
        jc = OxmlElement("w:lvlJc")
        jc.set(qn("w:val"), "left")
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "720")
        tabs.append(tab)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "720")
        ind.set(qn("w:hanging"), "360")
        p_pr.extend([tabs, ind])
        lvl.extend([start, fmt, text, jc, p_pr])
        if font:
            r_pr = OxmlElement("w:rPr")
            r_fonts = OxmlElement("w:rFonts")
            r_fonts.set(qn("w:ascii"), font)
            r_fonts.set(qn("w:hAnsi"), font)
            r_pr.append(r_fonts)
            lvl.append(r_pr)
        abstract.append(lvl)
        numbering.append(abstract)

    def make_num(num_id: int, abs_id: int) -> None:
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abs_ref = OxmlElement("w:abstractNumId")
        abs_ref.set(qn("w:val"), str(abs_id))
        num.append(abs_ref)
        numbering.append(num)

    make_abstract(next_abs, "bullet", "•", BODY_FONT)
    make_num(next_num, next_abs)
    make_abstract(next_abs + 1, "decimal", "%1.", BODY_FONT)
    make_num(next_num + 1, next_abs + 1)
    return next_num, next_num + 1


def apply_num(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])
    p_pr.append(num_pr)


def set_paragraph_shading(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def add_body(doc, text: str, *, bold_lead: str | None = None, color=TEXT, after=6, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    p.paragraph_format.keep_together = keep
    if bold_lead and text.startswith(bold_lead):
        add_run(p, bold_lead, bold=True, color=NAVY)
        add_run(p, text[len(bold_lead):], color=color)
    else:
        add_run(p, text, color=color)
    return p


def add_bullet(doc, text: str, bullet_id: int, *, bold_lead: str | None = None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    apply_num(p, bullet_id)
    if bold_lead and text.startswith(bold_lead):
        add_run(p, bold_lead, bold=True, color=NAVY)
        add_run(p, text[len(bold_lead):], color=TEXT)
    else:
        add_run(p, text, color=TEXT)
    return p


def add_numbered(doc, text: str, number_id: int, *, bold_lead: str | None = None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    apply_num(p, number_id)
    if bold_lead and text.startswith(bold_lead):
        add_run(p, bold_lead, bold=True, color=NAVY)
        add_run(p, text[len(bold_lead):], color=TEXT)
    else:
        add_run(p, text, color=TEXT)
    return p


def add_heading(doc, text: str, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    add_run(p, text, bold=True, color=ACCENT if level < 3 else NAVY, size={1: 16, 2: 13, 3: 12}[level])
    return p


def add_section_tag(doc, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    add_run(p, text.upper(), bold=True, color=GOLD, size=8.5)
    return p


def add_callout(doc, title: str, text: str, fill=LIGHT, accent=ACCENT):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_width(table)
    table.columns[0].width = Inches(0.12)
    table.columns[1].width = Inches(6.23)
    set_cell_width(table.cell(0, 0), 180)
    set_cell_width(table.cell(0, 1), 8980)
    set_cell_shading(table.cell(0, 0), accent)
    set_cell_shading(table.cell(0, 1), fill)
    set_cell_margins(table.cell(0, 0), 0, 0, 0, 0)
    cell = table.cell(0, 1)
    set_cell_margins(cell, 160, 180, 160, 180)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    add_run(p, title, bold=True, color=NAVY, size=11)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.08
    add_run(p2, text, color=TEXT, size=10)
    prevent_row_split(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_phase(doc, phase: str, title: str, problem: str, decision: str, result: str):
    tag = doc.add_paragraph()
    tag.paragraph_format.space_after = Pt(2)
    tag.paragraph_format.keep_with_next = True
    add_run(tag, phase, bold=True, color=GOLD, size=8.5)
    h = doc.add_paragraph()
    h.paragraph_format.space_after = Pt(5)
    h.paragraph_format.keep_with_next = True
    add_run(h, title, bold=True, color=NAVY, size=12)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.09
    add_run(p, "问题：", bold=True, color=ACCENT, size=10)
    add_run(p, problem + "  ", color=TEXT, size=10)
    add_run(p, "决定：", bold=True, color=ACCENT, size=10)
    add_run(p, decision + "  ", color=TEXT, size=10)
    add_run(p, "结果：", bold=True, color=ACCENT, size=10)
    add_run(p, result, color=TEXT, size=10)


def add_page_break(doc):
    # Major sections flow continuously. Heading keep rules provide clean
    # transitions without producing blank pages when a prior section already
    # reaches a natural page boundary.
    return None


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), CJK_FONT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, before, after, color in (
        ("Title", 30, 0, 12, NAVY),
        ("Subtitle", 14, 0, 16, MUTED),
        ("Heading 1", 16, 16, 8, ACCENT),
        ("Heading 2", 13, 12, 6, ACCENT),
        ("Heading 3", 12, 8, 4, NAVY),
    ):
        s = styles[style_name]
        s.font.name = BODY_FONT
        s.font.size = Pt(size)
        s.font.bold = style_name != "Subtitle"
        s.font.color.rgb = RGBColor.from_string(color)
        s._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), CJK_FONT)
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = style_name.startswith("Heading")


def configure_section(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)


def add_header_footer(section) -> None:
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    add_run(p, "SIGNAL ATLAS  ·  PRODUCT RECORD", bold=True, color=MUTED, size=8)

    footer = section.footer
    p = footer.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    table = footer.add_table(rows=1, cols=3, width=Inches(6.5))
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [3600, 3600, 2160]
    labels = ["产品记录 · 截至 2026-08-13", "内部复盘与后续规划", ""]
    aligns = [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT]
    for idx, cell in enumerate(table.rows[0].cells):
        set_cell_width(cell, widths[idx])
        set_cell_margins(cell, 0, 0, 0, 0)
        cell.text = ""
        cp = cell.paragraphs[0]
        cp.alignment = aligns[idx]
        cp.paragraph_format.space_after = Pt(0)
        if idx < 2:
            add_run(cp, labels[idx], color=MUTED, size=8)
        else:
            add_run(cp, "第 ", color=MUTED, size=8)
            add_page_field(cp)
            add_run(cp, " 页", color=MUTED, size=8)


def make_table(doc, headers, rows, widths, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(table)
    set_table_borders(table)
    hdr = table.rows[0]
    for i, text in enumerate(headers):
        set_cell_width(hdr.cells[i], widths[i])
        set_cell_shading(hdr.cells[i], "F2F4F7")
        set_cell_text(hdr.cells[i], text, bold=True, color=NAVY, size=9)
    set_repeat_header(table)
    prevent_row_split(hdr)
    for r_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, text in enumerate(row):
            set_cell_width(cells[i], widths[i])
            set_cell_text(cells[i], text, size=font_size)
            if r_idx % 2:
                set_cell_shading(cells[i], "FBFCFE")
        prevent_row_split(table.rows[-1])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def build() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    configure_section(doc.sections[0])
    add_header_footer(doc.sections[0])
    bullet_id, number_id = add_numbering_definitions(doc)

    # Cover / memo masthead
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    add_run(p, "PRODUCT RECORD  ·  0 → 1", bold=True, color=GOLD, size=9)

    p = doc.add_paragraph(style="Title")
    p.paragraph_format.space_after = Pt(6)
    add_run(p, "Signal Atlas", bold=True, color=NAVY, size=30)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(12)
    add_run(p2, "从 0 到 1 产品记录", bold=True, color=ACCENT, size=22)

    sub = doc.add_paragraph(style="Subtitle")
    add_run(sub, "从一次跨地区流量异常，到一个可共享的品牌舆情监测工具", color=MUTED, size=14)
    add_bottom_border(sub, color=ACCENT, size="16", space="8")

    meta = doc.add_table(rows=2, cols=2)
    meta.autofit = False
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(meta)
    set_table_borders(meta, color=WHITE, size="0")
    meta_rows = [
        ("记录范围", "从最初需求提出至当前公开版本（截至 2026 年 8 月 13 日）"),
        ("产品形态", "全球新闻与社交媒体舆情监测 Web 工具"),
        ("产品状态", "已公开发布，持续迭代"),
        ("公开地址", "https://signal-atlas-intelligence.rita-tanwenrui.chatgpt.site/"),
    ]
    flat = [(meta.cell(0, 0), meta_rows[0]), (meta.cell(0, 1), meta_rows[1]),
            (meta.cell(1, 0), meta_rows[2]), (meta.cell(1, 1), meta_rows[3])]
    for cell, (label, value) in flat:
        set_cell_shading(cell, LIGHTER)
        set_cell_margins(cell, 130, 150, 130, 150)
        cell.text = ""
        cp = cell.paragraphs[0]
        cp.paragraph_format.space_after = Pt(2)
        add_run(cp, label.upper(), bold=True, color=GOLD, size=8)
        vp = cell.add_paragraph()
        vp.paragraph_format.space_after = Pt(0)
        add_run(vp, value, color=TEXT, size=9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    add_callout(
        doc,
        "一句话结果",
        "把“偶然发现某个地区突然出现大量访问”的被动排查，转化为“配置品牌后自动发现、归档、聚类、追踪传播并分析舆情”的持续工作流。",
    )

    add_section_tag(doc, "Document intent")
    add_heading(doc, "这份文档记录什么", 1)
    add_body(doc, "这不是一份重新包装的 PRD，也不是功能清单。它记录的是产品如何从一个真实业务问题出发，在不断使用、发现偏差、纠正理解和补齐能力的过程中，逐步形成今天的结构。")
    add_body(doc, "文档既保留已经完成的决定，也明确当前仍受数据权限、接口稳定性和算法置信度限制的部分，供团队复盘、交接和下一阶段规划使用。")

    add_page_break(doc)

    add_section_tag(doc, "Origin")
    add_heading(doc, "1. 起点：一次异常流量暴露出的信息盲区", 1)
    add_body(doc, "公司接受一家香港媒体采访后，官网后台显示台湾和泰国的访问量明显上升。进一步追查才发现，台湾、泰国等地区的网页媒体与 Instagram 新闻账号已经转发了相关内容。")
    add_body(doc, "真正的问题并不是“这一次有哪些媒体转载”，而是：如果没有异常流量提醒，团队可能根本不知道报道已经跨地区传播；即使知道，也需要人工逐个搜索、判断来源、整理链接和估计传播顺序。")

    add_callout(
        doc,
        "最初的产品机会",
        "为任意品牌建立一个持续运行的全球监测工具：自动发现新闻和社交媒体提及，区分来源地区，保留历史档案，并解释一条信息是如何扩散、公众如何回应的。",
        fill="FFF8E8",
        accent=GOLD,
    )

    add_heading(doc, "2. 产品目标如何形成", 1)
    add_body(doc, "需求一开始看似只是“自动搜新闻”，但很快扩展为一条完整的数据链：发现只是入口，后续还需要归档、去重、事件聚类、地区识别、传播链路、评论采集、舆情分析、团队共享与报告输出。")
    add_bullet(doc, "通用，而非定制。页面不固定出现某个品牌名；品牌只在用户配置后进入查询与分析。", bullet_id, bold_lead="通用，而非定制。")
    add_bullet(doc, "自动，而非依赖手工。系统应主动搜索；手动录入只作为遗漏补充和证据修正。", bullet_id, bold_lead="自动，而非依赖手工。")
    add_bullet(doc, "解释，而非堆积。新闻不能只罗列，需要形成事件、时间线、传播网络和可行动的舆情判断。", bullet_id, bold_lead="解释，而非堆积。")
    add_bullet(doc, "共享，而非个人本地状态。同事登录后应看到同一个品牌配置、档案和分析结果。", bullet_id, bold_lead="共享，而非个人本地状态。")
    add_bullet(doc, "诚实呈现数据状态。未采集、无权限、未返回、无公开评论和确实为零必须分开。", bullet_id, bold_lead="诚实呈现数据状态。")

    add_heading(doc, "3. 从需求到产品的范围", 1)
    make_table(
        doc,
        ["层级", "要解决的问题", "最终形成的能力"],
        [
            ("发现", "外部报道发生了什么", "新闻源、免费爬虫、社交平台连接器与 Monid 发现"),
            ("组织", "资料如何长期保存", "统一档案、去重、来源地区、平台、互动信息与手动补录"),
            ("理解", "多篇内容是否属于同一事件", "基于时间连续性与标题/正文相似度的事件聚类"),
            ("传播", "事件从哪里开始、如何扩散", "以最早公开报道为锚点的有向传播图与置信度"),
            ("舆情", "公众如何讨论和回应", "评论采集、主题/词频/词云、多层情绪与风险复核"),
            ("协作", "如何供团队持续使用", "共享工作区、成员权限、连接器配置、报告导出"),
        ],
        [1120, 3050, 5190],
    )

    add_page_break(doc)

    add_section_tag(doc, "Iteration path")
    add_heading(doc, "4. 从 0 到 1 的关键迭代路径", 1)
    add_body(doc, "以下阶段不是预先排好的路线图，而是在真实使用中按问题出现的顺序逐步形成。每一轮都包含一次对产品边界或数据真实性的重新判断。")

    add_phase(doc, "阶段 01", "从 PRD 到可操作原型", "最初只有跨地区转载的案例与后台流量截图，缺少统一产品结构。", "先定义品牌配置、数据来源、地区识别、档案、分析与前端查看入口，并制作可操作的 Web 原型。", "把抽象需求变成可以点击、配置和验证的信息架构。")
    add_phase(doc, "阶段 02", "从案例展示转向通用产品", "早期页面把截图中的品牌和报道当作固定内容，容易误解为给单一公司制作的展示站。", "确认截图只是业务背景；移除 UI 中未配置就出现的品牌名与广告式文案，所有品牌信息改为用户输入后呈现。", "产品从一次性定制页面转为任意品牌都能使用的工作台。")
    add_phase(doc, "阶段 03", "自动新闻发现与混合采集架构", "用户不愿手工录入；GDELT 出现 HTTP 429、退避和零结果，单一免费来源无法稳定支撑。", "把监测窗口聚焦到近 31 天；采用少量全球发现服务、免费媒体爬虫和持续跟踪相结合的架构，并加入连接器配置入口。", "新闻发现从一次请求改为多源、可退避、可持续的管线。")
    add_phase(doc, "阶段 04", "新闻档案与基础舆情分析", "搜索结果只是堆叠列表，无法形成可查证的品牌历史记录，也缺少地区和情绪解释。", "建立统一档案表，按发布时间、平台、原文、媒体/账号与可获得互动归档；增加地区识别、词频、分词停用词、词云和情绪分析。", "每条结果成为可过滤、可追溯、可参与分析的数据对象。")
    add_phase(doc, "阶段 05", "同一事件识别与传播链路", "单篇文章分析无法解释一次报道如何在多个地区复制、改写并扩散。", "以时间连续性、标题与正文相似度聚类事件；以最早公开报道作为源点，根据发布时间与文本重合推断有向边。", "传播链路从并列指标改为按事件查看的箭头网络。")

    add_page_break(doc)

    add_phase(doc, "阶段 06", "品牌消歧与排除规则", "同名公司会污染结果，单靠品牌名不能精准定位。", "在品牌配置中加入别名、产品词、官方域名/账号、必须包含词和排除词；词典支持增加与删除。", "查询范围可由品牌自己定义，含排除词的内容会同步从档案、分析与评论舆情中移除。")
    add_phase(doc, "阶段 07", "社交媒体连接器与 Monid", "Instagram、Facebook、TikTok、X、YouTube 等平台没有统一免费官方搜索接口；用户希望在没有官方 API 时仍能接入。", "增加各平台与 Monid 的前端配置；把采集顺序明确为：关键词发现帖子、保存 URL、解析媒体 ID、采集评论/回复、内部模型分析。", "社交数据从“平台图标”变成有状态、有步骤、可核查的连接器。")
    add_phase(doc, "阶段 08", "评论舆情成为独立产品面", "新闻和帖子下的讨论比报道本身更能反映接受度，但早期界面大量显示“无法采集”或“无公开评论”，并不代表真实情况。", "建立评论专页，包含热门帖子、评论明细、风险复核和采集进度；提升信息密度，删除低价值的活跃作者模块；细分情绪。", "系统开始区分内容传播与受众反馈，并暴露采集失败的真实原因。")
    add_phase(doc, "阶段 09", "地图、可视化与交互可信度", "粗略或球面地图不适合比较全球地区，小地区不可见，悬停也没有数据；中国大陆与香港混算。", "改为平面展开的精细世界地图，加入缩放、悬停提示，并将中国大陆、香港、台湾等地区分别统计。", "地理分布从装饰性图形变成可读的分析入口。")
    add_phase(doc, "阶段 10", "团队共享、公开发布与自动报告", "个人浏览器状态无法支持同事协作；PostHog 只负责分析，不能充当网站托管；管理层还需要可流转的报告。", "建立团队共享工作区与成员权限，集中保存品牌和凭证配置；公开部署网站；加入基于当前数据自动生成舆情报告的能力。", "产品从个人原型转向可被团队共同使用、可对外访问的 SaaS 形态。")

    add_heading(doc, "阶段总览", 2)
    make_table(
        doc,
        ["阶段", "核心变化", "产品能力的提升"],
        [
            ("01–02", "可操作原型 → 通用品牌工作台", "从单一案例转为可复用产品"),
            ("03–04", "自动发现 → 统一档案与分析", "从找链接转为形成数据资产"),
            ("05–06", "事件传播 → 品牌消歧", "从内容堆积转为可解释、可控"),
            ("07–08", "社媒连接 → 评论舆情", "从报道监测扩展到受众反馈"),
            ("09–10", "可信可视化 → 团队与报告", "从个人原型走向团队工具"),
        ],
        [1200, 3480, 4680],
    )

    add_page_break(doc)

    add_section_tag(doc, "Corrections")
    add_heading(doc, "5. 关键纠偏：产品是在“发现不对”之后长出来的", 1)
    add_body(doc, "这一路最有价值的部分，并不是功能数量，而是多次把“看起来可以”的结果改成“真实可用”的判断规则。")

    corrections = [
        ("截图是背景，不是数据", "最初截图用于说明需求，不能直接作为产品默认数据。产品必须通过连接器自动发现，演示数据也要明确标识。"),
        ("无结果不等于没有报道", "零结果可能来自 API 限流、关键词过严、时间窗口、网络失败或配额用尽。前端要展示原因和重试状态。"),
        ("未采集不等于无公开评论", "平台未返回、尚未解析 ID、没有权限、任务排队和页面确实没有评论是不同状态，不能统一写成“无法采集”。"),
        ("时间空窗会切断事件", "不能把所有相似品牌报道都聚成一个事件。一个明显的一周空窗应强烈降低同事件概率；8 月 8 日后的连续爆发才可归入同一事件。"),
        ("最早报道必须是传播源点", "如果聚类内最早公开报道来自香港，传播图就不能从马来西亚开始。来源节点由时间锚定，后续边再由内容相似度推断。"),
        ("地区未知也需要解释", "媒体地区不能只靠页面声明；可组合域名、媒体知识库、语言、站点联系信息和历史来源，并保留置信度与人工修正入口。"),
        ("可视化必须帮助判断", "传播图需要箭头、平台色彩和图例；地图需要平面展开、细边界、缩放、悬停信息与地区拆分，而不是装饰性球面。"),
        ("通用产品必须支持消歧", "品牌名可能被多个主体共用。官方域名、账号、产品名、行业词、必须包含词和排除词共同决定结果是否属于目标品牌。"),
    ]
    for title, text in corrections:
        add_bullet(doc, title + "。" + text, bullet_id, bold_lead=title + "。")

    add_callout(
        doc,
        "最重要的产品原则",
        "宁可把状态标记为“待验证”或“连接器未返回”，也不要把系统没拿到的数据表达成现实世界中不存在。舆情工具的可信度首先来自对未知的诚实。",
        fill="FDEEEE",
        accent=RED,
    )

    add_page_break(doc)

    add_section_tag(doc, "Current product")
    add_heading(doc, "6. 当前产品结构", 1)
    add_body(doc, "当前版本已经形成从配置、采集到分析和协作的完整骨架。各模块围绕同一份品牌工作区数据运行，而不是各自保存一套孤立结果。")
    make_table(
        doc,
        ["模块", "主要功能", "用户价值"],
        [
            ("总览", "监测状态、核心指标、事件摘要、平面世界热力地图", "快速判断报道规模、地区集中度和异常变化"),
            ("新闻与社媒档案", "统一表格、筛选、排序、原文链接、来源与互动", "保留长期、可核查的品牌媒体记录"),
            ("事件与传播链路", "事件聚类、最早源点、时间线、有向传播图、平台图例", "理解一次报道如何跨媒体、平台和地区扩散"),
            ("舆情分析", "主题、词频、词云、语言、地区与细分情绪", "从大量内容中提炼讨论焦点与态度结构"),
            ("评论舆情", "热门帖子、评论明细、风险队列、采集覆盖", "把受众反馈单独管理并支持人工复核"),
            ("来源覆盖", "新闻 API、免费爬虫、Monid 与各社媒连接器状态", "配置数据入口并定位失败环节"),
            ("品牌配置", "别名、产品词、官方资产、包含词、排除词、词典", "减少同名品牌和无关结果污染"),
            ("团队工作区", "共享品牌、档案、分析、成员与权限、凭证保险库", "同事登录后直接继承同一工作上下文"),
            ("报告", "按当前数据自动生成可下载的舆情分析报告", "把监测结果转为可汇报、可存档的业务材料"),
            ("手动补录", "补充系统遗漏的新闻或社媒链接并参与后续分析", "保留人工校正和应急补充能力"),
        ],
        [1550, 4300, 3510],
        font_size=8.8,
    )

    add_heading(doc, "当前公开入口", 2)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    add_run(p, "网站地址：", bold=True, color=NAVY)
    add_hyperlink(p, "signal-atlas-intelligence.rita-tanwenrui.chatgpt.site", "https://signal-atlas-intelligence.rita-tanwenrui.chatgpt.site/")
    add_body(doc, "公开发布解决的是“别人能打开网站”；团队共享工作区解决的是“别人登录后能看到同一份数据”。PostHog 仅用于网站行为分析，不承担托管和共享数据。")

    add_page_break(doc)

    add_section_tag(doc, "Data flow")
    add_heading(doc, "7. 核心数据链路", 1)
    add_heading(doc, "7.1 新闻链路", 2)
    news_steps = [
        "读取品牌名、别名、产品词、官方域名、包含词与排除词，生成多语言查询。",
        "通过全球发现服务寻找近期报道，再由免费媒体爬虫与已知来源库持续跟踪。",
        "规范 URL、标题、时间和正文，去重并保留原始证据。",
        "根据媒体知识库、域名、语言和站点信息推断来源地区，记录置信度。",
        "将合格内容写入统一档案；含排除词的结果不进入档案、分析与评论舆情。",
        "按时间连续性与文本相似度聚成事件，再计算最早源点、可能传播边和分析指标。",
    ]
    for item in news_steps:
        add_numbered(doc, item, number_id)

    add_heading(doc, "7.2 社交媒体与评论链路", 2)
    social_steps = [
        "Monid 或平台连接器根据普通关键词、品牌别名、产品词和账号范围搜索相关帖子。",
        "保存帖子 URL、平台、发布时间、作者 ID、公开粉丝数和可见互动数据。",
        "由 URL 解析平台所需的媒体 ID、视频 ID 或帖子 ID；解析失败单独标记并重试。",
        "调用可用的评论/回复采集能力；平台未返回与确实无评论必须分开记录。",
        "对评论做语言识别、分词、停用词过滤、主题归纳、细分情绪和风险标记。",
        "将帖子和评论关联到品牌、事件与工作区，在评论舆情页集中展示。",
    ]
    for item in social_steps:
        add_numbered(doc, item, number_id)

    add_callout(
        doc,
        "正确的顺序",
        "关键词搜帖子 → 保存帖子 URL → 解析帖子标识 → 采集评论/回复 → 网站内部分析。手动粘贴帖子 URL 只是补充入口，不应代替自动发现。",
    )

    add_page_break(doc)

    add_section_tag(doc, "Product logic")
    add_heading(doc, "8. 三个关键分析模型", 1)
    add_heading(doc, "8.1 品牌归属判断", 2)
    add_body(doc, "品牌相关性不是“正文里出现了品牌名”这么简单。系统需要同时考虑正向证据和排除证据，并允许团队维护词典。")
    add_bullet(doc, "正向证据：官方域名/账号、产品名、核心人物、行业词、品牌别名、必须包含词。", bullet_id, bold_lead="正向证据：")
    add_bullet(doc, "负向证据：同名主体、其他行业语境、排除词、非目标地区或语言组合。", bullet_id, bold_lead="负向证据：")
    add_bullet(doc, "输出：相关 / 待复核 / 排除，并保留触发判断的证据。", bullet_id, bold_lead="输出：")

    add_heading(doc, "8.2 同一事件聚类", 2)
    add_body(doc, "事件聚类同时看时间和内容。标题与正文高度相似只能说明内容接近；如果发布时间之间存在明显空窗，就不应轻易合并。反过来，同一波传播也可能因为翻译、改写而出现标题差异，需要正文实体与共同事实补充判断。")
    add_bullet(doc, "时间信号：发布间隔、爆发窗口、空窗长度。", bullet_id, bold_lead="时间信号：")
    add_bullet(doc, "文本信号：标题重合、正文相似、关键实体与事实组合。", bullet_id, bold_lead="文本信号：")
    add_bullet(doc, "人工控制：允许拆分、合并和确认事件，修正后影响传播链路与报告。", bullet_id, bold_lead="人工控制：")

    add_heading(doc, "8.3 传播路径推断", 2)
    add_body(doc, "传播图是一种基于公开时间与文本相似度的推断，不等同于平台提供的真实转发日志。最早公开报道作为根节点；后续节点只能指向比自己更早、且内容更接近的候选来源。")
    add_bullet(doc, "节点：媒体或账号发布的一条内容；颜色用于区分网页新闻、Instagram、Facebook、TikTok、X 与 YouTube。", bullet_id, bold_lead="节点：")
    add_bullet(doc, "边：可能的引用、改写或转载关系；箭头方向必须满足时间先后。", bullet_id, bold_lead="边：")
    add_bullet(doc, "置信度：由时间差、标题/正文相似、引用线索、来源影响力与地区跳转共同决定。", bullet_id, bold_lead="置信度：")

    add_page_break(doc)

    add_section_tag(doc, "Design principles")
    add_heading(doc, "9. 迭代中形成的产品原则", 1)
    principles = [
        ("真实数据优先于漂亮仪表盘", "先确认数据从哪里来、为什么缺失，再讨论视觉呈现。演示数据必须明确标识，不能让人误以为是真实监测结果。"),
        ("自动化是主流程，人工是纠错机制", "工具的价值在于持续发现；手动补录、地区修正、事件拆分和情绪复核用于处理边缘情况。"),
        ("状态需要可解释", "连接器应告诉用户是在排队、限流、缺凭证、解析失败、未返回还是确实为空，并给出下一步。"),
        ("分析必须回到证据", "情绪、事件和传播边都应能追溯到原文、评论、时间戳和判断依据。"),
        ("共享的是工作区，不只是网址", "公开网址使别人能访问；账户、权限、数据库和共享配置才让团队看到同一份内容。"),
        ("简洁不是减少信息，而是降低误解", "去掉广告式标语和重复说明，把空间留给数据、证据、状态和操作。"),
    ]
    for title, text in principles:
        add_bullet(doc, title + "。" + text, bullet_id, bold_lead=title + "。")

    add_heading(doc, "10. 当前边界与风险", 1)
    add_body(doc, "当前产品已经具备完整工作流，但仍有几类限制不能通过前端文案掩盖。")
    risks = [
        ("数据权限", "只能采集公开、可访问的内容；私密、删除、登录受限或平台禁止抓取的数据不能保证获得。"),
        ("连接器稳定性", "Monid、新闻 API 和免费爬虫都可能受配额、限流、页面变更和供应商能力影响，需要队列、退避与监控。"),
        ("评论覆盖", "帖子发现成功不代表评论采集成功。不同平台需要不同 ID 解析和权限，覆盖率应逐平台统计。"),
        ("地区推断", "来源地区可通过多种信号推断，但跨国媒体、镜像站和聚合站仍可能需要人工修正。"),
        ("传播因果", "公开数据只能推断可能路径，不能把文本相似和时间先后表述为确定的转发关系。"),
        ("情绪语境", "讽刺、多语言俚语、成人话题和文化差异会影响自动分类，需要保留复核队列与证据。"),
        ("合规", "社交平台采集必须遵守服务条款、隐私要求和适用地区法规，凭证需要加密保存并限制权限。"),
    ]
    make_table(doc, ["风险面", "当前边界与应对"], risks, [1700, 7660], font_size=9.0)

    add_page_break(doc)

    add_section_tag(doc, "Learning")
    add_heading(doc, "11. 这次从 0 到 1 得到的经验", 1)
    add_heading(doc, "产品层面", 2)
    add_bullet(doc, "用户真正要的不是“能搜”，而是从发现到判断、从判断到协作的闭环。", bullet_id)
    add_bullet(doc, "一个看似很小的异常流量案例，可以抽象成多品牌、多地区都成立的通用问题。", bullet_id)
    add_bullet(doc, "最有价值的迭代通常来自用户指出“这个结论明显不对”，而不是继续增加页面。", bullet_id)

    add_heading(doc, "数据与工程层面", 2)
    add_bullet(doc, "全球舆情监测无法依赖单一免费 API；混合来源、缓存、去重、退避和可观测性是基础设施。", bullet_id)
    add_bullet(doc, "社交媒体的帖子发现、详情获取和评论采集是三类不同任务，不能用一个“已连接”状态概括。", bullet_id)
    add_bullet(doc, "任何推断模型都要把时间约束、置信度和人工修正作为一等数据。", bullet_id)

    add_heading(doc, "体验层面", 2)
    add_bullet(doc, "工具文案应描述事实和下一步，不使用生硬的营销口号。", bullet_id)
    add_bullet(doc, "信息密度应与任务匹配：评论复核适合密集卡片，传播关系适合一张主图，档案适合表格。", bullet_id)
    add_bullet(doc, "小地区、未知地区和无结果状态都是核心体验，不能只在“有数据时”设计。", bullet_id)

    add_heading(doc, "12. 下一阶段路线图", 1)
    make_table(
        doc,
        ["优先级", "方向", "完成标准"],
        [
            ("P0", "采集可靠性与可观测性", "每个连接器显示成功率、延迟、限流、失败原因、重试和最近有效数据；不再用模糊空状态。"),
            ("P0", "评论覆盖提升", "按平台打通 URL→ID→评论链路，输出真实覆盖率，并支持指定帖子补抓。"),
            ("P0", "来源地区质量", "建立媒体来源知识库、置信度与人工修正回写，减少“地区未披露”。"),
            ("P1", "事件聚类与传播校准", "加入可调时间窗口、空窗惩罚、人工拆并和传播边解释；用已知案例回归测试。"),
            ("P1", "多语言语义与情绪", "提升翻译改写、讽刺和细分情绪识别，保留证据句与复核反馈。"),
            ("P1", "预警与定期报告", "支持关键词/情绪/传播速度阈值告警，以及日报、周报、事件报告自动生成。"),
            ("P2", "竞品与基准", "允许在同一工作区比较多个品牌的声量、地区、情绪与事件响应。"),
            ("P2", "团队治理", "完善角色权限、审计日志、数据保留策略与连接器凭证轮换。"),
        ],
        [900, 2800, 5660],
        font_size=8.8,
    )

    add_page_break(doc)

    add_section_tag(doc, "Closing")
    add_heading(doc, "13. 结语：这个产品现在处在什么位置", 1)
    add_body(doc, "Signal Atlas 已经从一个针对单次跨地区转载的临时想法，演变成一个拥有完整骨架的品牌舆情监测产品：品牌可配置、数据可自动发现、历史可归档、事件可聚类、传播可推断、评论可分析、结果可共享和导出。")
    add_body(doc, "但它仍不是一个“数据已经覆盖一切”的成熟平台。下一阶段的核心不是继续增加更多看起来丰富的模块，而是把连接器成功率、评论覆盖、地区判断、事件边界和传播置信度做得更稳定、更透明，并用真实事件持续校准。")
    add_callout(
        doc,
        "从 0 到 1 的真正完成标志",
        "不是页面已经上线，而是团队遇到下一次跨地区传播时，不再从手工搜索开始；系统已经保存了品牌范围、持续采集证据、形成事件，并能解释“发生了什么、从哪里开始、如何扩散、公众怎样回应”。",
        fill="EDF7F2",
        accent=GREEN,
    )

    add_heading(doc, "附录：主要迭代清单", 2)
    changelog = [
        "建立通用品牌配置与自动监测入口，移除固定品牌文案。",
        "引入近 31 天新闻发现、免费爬虫持续追踪与 API 配置。",
        "建立统一新闻/社媒档案及手动补录。",
        "增加来源地区推断、平面缩放世界地图和地区拆分。",
        "增加停用词过滤、词频、词云、细分情绪与舆情分析。",
        "按时间空窗和内容相似度聚类同一事件。",
        "以最早公开报道为源点绘制带平台颜色的传播图。",
        "增加 Monid 与五类社媒连接器，形成帖子发现、URL/ID、评论采集和分析链路。",
        "增加评论舆情专页、真实采集状态、品牌消歧、全局排除词与词典删除。",
        "建立团队共享工作区、成员权限、公开发布与自动报告。",
    ]
    change_table = doc.add_table(rows=5, cols=2)
    change_table.autofit = False
    change_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(change_table)
    for row_idx, row in enumerate(change_table.rows):
        prevent_row_split(row)
        for col_idx, cell in enumerate(row.cells):
            set_cell_width(cell, 4680)
            set_cell_margins(cell, 70, 80, 70, 80)
            if row_idx % 2 == 0:
                set_cell_shading(cell, LIGHTER)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            apply_num(p, bullet_id)
            add_run(p, changelog[col_idx * 5 + row_idx], color=TEXT, size=9.2)

    # Core metadata
    props = doc.core_properties
    props.title = "Signal Atlas 从 0 到 1 产品记录"
    props.subject = "品牌舆情监测产品的需求起点、关键迭代、纠偏、当前结构与路线图"
    props.author = "Signal Atlas 项目团队"
    props.keywords = "Signal Atlas, 产品记录, 舆情监测, 新闻监测, 社交媒体, 产品迭代"
    props.comments = "基于产品从需求提出到当前公开版本的真实迭代过程整理。"

    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    build()
