from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path("output/Signal-Atlas-产品迭代需求池.docx")

INK = "17233C"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
TEAL = "167D86"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHT_TEAL = "E8F4F4"
LIGHT_GRAY = "F2F4F7"
PALE_GOLD = "FFF7E0"
GOLD = "7A5A00"
PALE_RED = "FDECEC"
RED = "9B1C1C"
WHITE = "FFFFFF"
BORDER = "C8D2DC"


def set_run_font(run, size=11, bold=False, color=INK, italic=False):
    run.font.name = "Arial Unicode MS"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), "Arial Unicode MS")
    rfonts.set(qn("w:hAnsi"), "Arial Unicode MS")
    rfonts.set(qn("w:eastAsia"), "Arial Unicode MS")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = tc_borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tc_borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[idx]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            set_cell_margins(cell)
            set_cell_borders(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def style_cell_text(cell, size=9.2, bold=False, color=INK, alignment=WD_ALIGN_PARAGRAPH.LEFT):
    for p in cell.paragraphs:
        p.alignment = alignment
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        for run in p.runs:
            set_run_font(run, size=size, bold=bold, color=color)


def add_table(doc, headers, rows, widths_dxa, aligns=None, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = 0
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        set_cell_shading(cell, LIGHT_BLUE)
        style_cell_text(cell, size=9.2, bold=True, color=DARK_BLUE,
                        alignment=(aligns[i] if aligns else WD_ALIGN_PARAGRAPH.LEFT))
    set_repeat_table_header(table.rows[0])
    for row_data in rows:
        row = table.add_row()
        for i, value in enumerate(row_data):
            row.cells[i].text = str(value)
            style_cell_text(row.cells[i], size=font_size,
                            alignment=(aligns[i] if aligns else WD_ALIGN_PARAGRAPH.LEFT))
    set_table_geometry(table, widths_dxa)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    return table


def add_para(doc, text="", size=11, color=INK, bold=False, italic=False,
             before=0, after=6, align=WD_ALIGN_PARAGRAPH.LEFT, keep=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.keep_with_next = keep
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color, italic=italic)
    return p


def add_labeled_para(doc, label, text, after=5):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    r1 = p.add_run(label)
    set_run_font(r1, size=10.5, bold=True, color=DARK_BLUE)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.5, color=INK)
    return p


def add_heading(doc, text, level=1):
    style = doc.styles[f"Heading {level}"]
    p = doc.add_paragraph(style=style)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_run_font(
        run,
        size={1: 16, 2: 13, 3: 12}[level],
        bold=True,
        color={1: BLUE, 2: BLUE, 3: DARK_BLUE}[level],
    )
    return p


def add_callout(doc, label, text, fill=LIGHT_TEAL, label_color=TEAL):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_borders(cell, color=fill, size="2")
    set_cell_margins(cell, top=140, start=180, bottom=140, end=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.18
    r1 = p.add_run(label + "  ")
    set_run_font(r1, size=10.5, bold=True, color=label_color)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.5, color=INK)
    set_table_geometry(table, [9360], indent_dxa=120)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def set_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    run2 = paragraph.add_run(" 页")
    set_run_font(run2, size=9, color=MUTED)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial Unicode MS"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial Unicode MS")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial Unicode MS")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        1: (16, BLUE, 18, 10),
        2: (13, BLUE, 14, 7),
        3: (12, DARK_BLUE, 10, 5),
    }
    for level, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Arial Unicode MS"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial Unicode MS")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial Unicode MS")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    hr = hp.add_run("SIGNAL ATLAS  /  产品迭代需求池")
    set_run_font(hr, size=9, bold=True, color=MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    set_page_number(fp)


def build_document():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)

    add_para(doc, "产品迭代需求池", size=26, bold=True, color=INK, after=4, keep=True)
    add_para(doc, "Signal Atlas｜持续收集、统一整理、分批实施", size=13, color=TEAL, after=14, keep=True)

    metadata = [
        ("文档状态", "持续收集中（已完成项单独标记）"),
        ("版本", "v1.2"),
        ("建立日期", "2026年8月13日"),
        ("实施触发", "仅在明确收到“整理这一批并开始修改”后启动"),
    ]
    meta = doc.add_table(rows=0, cols=2)
    meta.style = "Table Grid"
    for label, value in metadata:
        row = meta.add_row()
        row.cells[0].text = label
        row.cells[1].text = value
        set_cell_shading(row.cells[0], LIGHT_BLUE)
        style_cell_text(row.cells[0], size=9.5, bold=True, color=DARK_BLUE)
        style_cell_text(row.cells[1], size=9.5)
    set_table_geometry(meta, [1800, 7560])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    add_callout(
        doc,
        "工作约定",
        "以后新增的单条需求先进入本文件，不立即改代码。我会保留原始表述，再完成归类、去重、依赖分析、优先级排序和验收标准；形成一批后，先提交统一修改策略，再集中实施和验收。",
    )

    add_heading(doc, "1. 文档用途与工作规则", 1)
    add_labeled_para(doc, "唯一需求来源：", "本文件作为后续迭代的统一需求池，避免需求散落在多轮对话中。")
    add_labeled_para(doc, "保留原话：", "每条需求同时保存用户原始表述和整理后的产品要求，减少理解偏差。")
    add_labeled_para(doc, "默认不实施：", "在需求收集阶段只记录和分析，不因为新增一条意见就立即修改网站。")
    add_labeled_para(doc, "集中实施：", "收到明确启动指令后冻结本轮范围，按依赖关系集中修改、测试、发布。")
    add_labeled_para(doc, "可追溯：", "每项需求均有编号、状态、优先级、验收标准、实施批次和变更记录。")

    add_heading(doc, "2. 从提出需求到发布的统一流程", 1)
    add_table(
        doc,
        ["阶段", "处理内容", "产出", "是否改代码"],
        [
            ("收集", "记录原始表述、日期和使用场景", "待整理需求", "否"),
            ("归类", "按产品模块归档，识别重复或冲突", "需求分类", "否"),
            ("澄清", "补齐目标、范围、数据来源和边界", "可执行需求", "否"),
            ("排序", "评估影响、紧急度、依赖和实施成本", "优先级与顺序", "否"),
            ("策略", "确定数据、算法、接口、页面和迁移方案", "统一修改方案", "否"),
            ("实施", "按冻结范围集中开发和数据迁移", "候选版本", "是"),
            ("验收", "按验收标准测试、修正并发布", "发布版本与记录", "是"),
        ],
        [1200, 3850, 2710, 1600],
        aligns=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT,
                WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER],
    )

    add_heading(doc, "3. 分类、状态与优先级标准", 1)
    add_heading(doc, "3.1 产品模块分类", 2)
    add_table(
        doc,
        ["模块", "包含内容"],
        [
            ("数据采集", "新闻、社媒帖子、评论、回复、互动数据与连接器"),
            ("品牌配置", "品牌实体、产品词、别名、排除词和地区范围"),
            ("情感与语义分析", "情绪、立场、意图、主题、风险、人工校准与模型版本"),
            ("事件与传播", "事件聚类、转载判断、跨地区传播路径与证据"),
            ("评论舆情", "评论档案、互动加权、话题聚类、风险复核和结论"),
            ("可视化与报告", "地图、趋势图、对比分析、导出和自动报告"),
            ("团队与权限", "共享工作区、账号、权限、审计和配置共享"),
            ("系统质量", "性能、稳定性、数据完整性、可解释性和测试"),
        ],
        [2300, 7060],
    )

    add_heading(doc, "3.2 状态定义", 2)
    add_table(
        doc,
        ["状态", "含义"],
        [
            ("待整理", "已记录原始需求，尚未完成结构化分析"),
            ("待澄清", "存在影响范围或实现路径不明确的问题"),
            ("已归类", "需求结构完整，等待进入实施批次"),
            ("本轮计划", "已冻结到下一实施批次"),
            ("实施中", "正在开发、迁移或测试"),
            ("待验收", "已完成实现，等待按标准检查"),
            ("已完成", "验收通过并已发布"),
            ("延后", "暂不进入近期实施，但保留记录"),
        ],
        [1900, 7460],
    )

    add_heading(doc, "3.3 优先级", 2)
    add_table(
        doc,
        ["级别", "判断标准"],
        [
            ("P0", "阻断采集、导致数据丢失、权限或安全风险；需优先处理"),
            ("P1", "核心分析明显错误，直接影响产品可信度或主要工作流程"),
            ("P2", "显著提升效率和分析价值，但不阻断核心使用"),
            ("P3", "体验优化、视觉改进或长期能力储备"),
        ],
        [1900, 7460],
    )

    doc.add_page_break()
    add_heading(doc, "4. 当前需求池总览", 1)
    add_callout(
        doc,
        "当前状态",
        "本批需求已完成收集、归类和初步排序。SA-REQ-005 已完成开发与发布；其余需求仍等待统一冻结实施范围。建议下一步先完成数据保护，再建设多语言判断、互动加权和分析 Agent。",
        fill=PALE_GOLD,
        label_color=GOLD,
    )
    add_table(
        doc,
        ["编号", "需求", "模块", "优先级", "状态"],
        [
            ("SA-REQ-001", "接入成熟的多语言情感模型", "情感与语义", "P1", "已归类"),
            ("SA-REQ-002", "保留人工标注并校准模型", "情感与语义", "P1", "已归类"),
            ("SA-REQ-003", "评论舆情分析 Agent 与丰富指标", "评论舆情", "P1", "已归类"),
            ("SA-REQ-004", "按评论获赞与回复数进行加权", "评论舆情", "P1", "已归类"),
            ("SA-REQ-005", "翻译能力与 Monid 解耦", "数据采集", "P1", "已完成"),
        ],
        [1600, 3900, 1700, 1000, 1160],
        aligns=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT,
                WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER,
                WD_ALIGN_PARAGRAPH.CENTER],
        font_size=8.8,
    )

    doc.add_page_break()
    add_heading(doc, "5. 已收集需求明细", 1)

    add_heading(doc, "SA-REQ-001｜接入成熟的多语言情感模型", 2)
    add_labeled_para(doc, "原始需求：", "网上已有较成熟、面向不同语言和国家的情感分析算法，希望网站使用更可靠的能力。")
    add_labeled_para(doc, "整理后要求：", "引入多语言语义模型作为基础判断层，支持简体中文、繁体中文、英语、泰语、日语、韩语及其他主要市场语言，并保留语言和地区维度的独立评估。")
    add_labeled_para(doc, "范围边界：", "国家差异不能仅靠语言代码推断；同一种语言在不同地区的俚语、反讽和文化语境需要分别校准。")
    add_labeled_para(doc, "验收标准：", "模型结果包含标签、分数、置信度、模型版本和语言；可按语言/地区查看准确率；低置信度评论进入复核队列。")
    add_labeled_para(doc, "依赖：", "依赖SA-REQ-002的标注保留与模型版本机制。")

    add_heading(doc, "SA-REQ-002｜保留人工标注并校准模型", 2)
    add_labeled_para(doc, "原始需求：", "已经人工标注了很多评论；模型更新后必须保留这些标注，不希望重新从头标注。")
    add_labeled_para(doc, "整理后要求：", "人工标注作为不可覆盖的金标准独立保存；新增模型预测采用版本化记录，允许比较旧模型、新模型和人工结果。")
    add_labeled_para(doc, "迁移要求：", "升级前生成标注快照和数量校验；任何再分析操作不得修改人工标签。")
    add_labeled_para(doc, "验收标准：", "升级前后人工标注数量、内容、人员和时间保持一致；页面始终优先显示人工判断；支持回滚模型版本。")
    add_labeled_para(doc, "依赖：", "是其他情感分析升级的基础能力，应最先实施。")

    add_heading(doc, "SA-REQ-003｜评论舆情分析 Agent 与丰富指标", 2)
    add_labeled_para(doc, "原始需求：", "现有评论分析过于浅显，希望Agent结合品牌情况自动给出有实质意义的分析，并增加更丰富、实用的指标。")
    add_labeled_para(doc, "整理后要求：", "先完成逐条结构化判断和聚类，再由品牌分析Agent读取品牌资料、指标、代表性评论和异常信号，形成带证据的结论与行动建议。")
    add_labeled_para(doc, "核心输出：", "主要议题、情绪原因、购买意向、质疑与投诉、地区/平台差异、高互动风险、误解与谣言、趋势变化、建议回应以及结论置信度。")
    add_labeled_para(doc, "验收标准：", "每条关键结论可追溯到帖子和代表性评论；区分事实、推断和建议；数据不足时明确提示，不生成空泛结论。")
    add_labeled_para(doc, "依赖：", "依赖SA-REQ-001的稳定分类结果和SA-REQ-004的加权指标。")

    add_heading(doc, "SA-REQ-004｜评论互动加权", 2)
    add_labeled_para(doc, "原始需求：", "希望评论或赞数按照加权方式计分，获赞越多，情绪权重越高，而不是只按照评论数量统计。")
    add_labeled_para(doc, "整理后要求：", "保留数量口径，同时增加互动加权口径。评论层使用评论本身的获赞数和回复数；帖子层单独使用点赞、评论、分享和播放量，不重复放大。")
    add_labeled_para(doc, "计算原则：", "采用对数降权、平台内标准化和权重上限；Instagram、YouTube、TikTok等平台分别比较，缺失互动数据不得伪装成0。")
    add_labeled_para(doc, "验收标准：", "页面同时显示普通情绪占比、互动加权占比及差值；能够解释权重来源；单条爆款评论不会无限压过其他观点。")
    add_labeled_para(doc, "依赖：", "需要采集到可靠的评论获赞数、回复数和平台字段。")

    add_heading(doc, "SA-REQ-005｜翻译能力与 Monid 解耦", 2)
    add_labeled_para(doc, "实施状态：", "已于2026-08-13完成开发并发布。Monid只负责社媒数据采集，翻译由独立队列处理；新闻档案优先翻译，简体中文、繁体中文和英语自动跳过。")
    add_labeled_para(doc, "原始需求：", "翻译不需要接 Monid，任何工具都可以做到。")
    add_labeled_para(doc, "问题现状：", "当前通用翻译任务被错误地放在 Monid 社媒采集流程内；当 Monid 未配置、未到运行时间或连接器异常时，网页新闻和其他来源的翻译也可能停滞，并在前端长期显示笼统的自动重试提示。")
    add_labeled_para(doc, "整理后要求：", "将翻译建设为与采集连接器无关的共享服务。网页新闻、社媒帖子、评论和回复在写入或内容更新后均可独立进入翻译队列，Monid 只负责采集，不负责触发或提供翻译。")
    add_labeled_para(doc, "语言规则：", "入队前先做语言识别；简体中文、繁体中文和英语原文不翻译，其他语言翻译为英文。语言缺失或置信度不足时重新识别，不因错误元数据盲目翻译。")
    add_labeled_para(doc, "技术策略：", "使用可替换的翻译提供商适配层，可接独立翻译 API、通用语言模型或自托管模型；提供商故障时不阻断采集和分析，并支持备用提供商、指数退避和失败重试。")
    add_labeled_para(doc, "状态与数据：", "保存来源文本哈希、检测语言、提供商、任务状态、尝试次数、最后错误、下次重试时间和完成时间；前端显示排队中、翻译中、失败原因及预计重试时间，不再统一显示“Translation will retry automatically”。")
    add_labeled_para(doc, "验收标准：", "未配置 Monid 时，非中英文内容仍能正常翻译；中英文原文不进入翻译队列；任一翻译提供商异常时可查看真实错误和后续动作；采集任务不因翻译失败而失败。")
    add_labeled_para(doc, "依赖：", "是 SA-REQ-001 多语言分析前的数据基础，可独立于社媒连接器实施。")

    add_heading(doc, "6. 建议修改顺序与策略", 1)
    add_table(
        doc,
        ["顺序", "实施批次", "主要工作", "完成标志"],
        [
            ("1", "数据保护与版本化", "备份人工标注；建立模型版本、预测结果和回滚机制", "旧标注零丢失，可追溯"),
            ("2", "独立翻译基础", "语言识别、独立队列、提供商适配、失败与重试状态", "不配置Monid也可翻译"),
            ("3", "多语言判断层", "基础模型、地区校准、低置信度复核", "替代词典规则作为主判断"),
            ("4", "互动加权层", "平台内标准化、评论/帖子分层权重、双口径指标", "普通与加权结果并列"),
            ("5", "评论分析 Agent", "主题聚类、异常检测、证据提取和品牌结论", "结论可追溯、可复核"),
            ("6", "页面与报告", "重构评论舆情页面并同步报告导出", "页面、CSV、报告口径一致"),
        ],
        [850, 2050, 4210, 2250],
        aligns=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT,
                WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
        font_size=8.9,
    )

    add_heading(doc, "7. 集中实施的启动与冻结规则", 1)
    add_labeled_para(doc, "默认行为：", "你继续提出需求时，我只更新需求池，并简要反馈新增编号、归类和可能依赖。")
    add_labeled_para(doc, "启动指令：", "当你说“整理这一批并开始修改”“本轮开始实施”或表达同等含义时，进入实施准备。")
    add_labeled_para(doc, "实施前输出：", "先提交本轮范围、合并/冲突项、优先顺序、数据迁移、风险、验收标准和预计影响页面。")
    add_labeled_para(doc, "范围冻结：", "启动后新增需求默认进入下一批，除非属于P0阻断问题或你明确要求并入本轮。")
    add_labeled_para(doc, "完成标准：", "实现、数据校验、测试、发布和需求状态更新全部完成后，才标记为已完成。")

    add_heading(doc, "8. 新需求记录模板", 1)
    add_table(
        doc,
        ["字段", "记录内容"],
        [
            ("需求编号", "SA-REQ-XXX"),
            ("提出日期", "YYYY-MM-DD"),
            ("原始表述", "保留用户原话，不先行改写"),
            ("问题与目标", "当前问题、使用场景、期望结果"),
            ("产品模块", "按第3.1节分类"),
            ("优先级与理由", "P0/P1/P2/P3及判断依据"),
            ("依赖与冲突", "前置数据、接口、算法、页面或其他需求"),
            ("验收标准", "可观察、可测试、可判定是否完成"),
            ("实施批次", "待排期或具体批次"),
            ("状态", "按第3.2节更新"),
        ],
        [2100, 7260],
    )

    doc.add_page_break()
    add_heading(doc, "9. 变更记录", 1)
    add_table(
        doc,
        ["版本", "日期", "变更内容"],
        [
            ("v1.2", "2026-08-13", "完成SA-REQ-005并发布：翻译与Monid解耦，启用独立翻译队列、新闻档案优先、语言跳过、真实错误与退避重试。"),
            ("v1.1", "2026-08-13", "新增SA-REQ-005：翻译能力与Monid解耦；记录独立队列、语言跳过、可替换提供商和透明重试要求。"),
            ("v1.0", "2026-08-13", "建立统一需求池；收录多语言情感模型、人工标注保留、评论分析Agent和互动加权需求。"),
        ],
        [1300, 1800, 6260],
    )

    # Prevent Word from splitting compact tables awkwardly where possible.
    for table in doc.tables:
        for row in table.rows:
            tr_pr = row._tr.get_or_add_trPr()
            cant_split = OxmlElement("w:cantSplit")
            tr_pr.append(cant_split)

    doc.core_properties.title = "Signal Atlas 产品迭代需求池"
    doc.core_properties.subject = "持续需求收集、归类、排序与分批实施"
    doc.core_properties.author = "Signal Atlas 产品团队"
    doc.core_properties.keywords = "Signal Atlas, 产品需求, 迭代, backlog"
    doc.save(OUTPUT)
    print(OUTPUT.resolve())


if __name__ == "__main__":
    build_document()
